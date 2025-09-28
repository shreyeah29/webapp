from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION


THEME_PRIMARY = RGBColor(38, 70, 83)       # Deep teal
THEME_SECONDARY = RGBColor(233, 196, 106)  # Warm sand
THEME_ACCENT = RGBColor(231, 111, 81)      # Coral
THEME_TEXT = RGBColor(33, 37, 41)          # Near-black
THEME_MUTED = RGBColor(136, 149, 161)      # Muted gray


def set_text_format(text_frame, font_size=24, bold=False, color=THEME_TEXT):
    for paragraph in text_frame.paragraphs:
        paragraph.font.size = Pt(font_size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color


def add_brand_bar(slide, title_text=None):
    left, top, width, height = Inches(0), Inches(0), Inches(13.33), Inches(0.4)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = THEME_PRIMARY
    bar.line.fill.background()

    if title_text:
        tx = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(12.5), Inches(0.3))
        tf = tx.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title_text
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)


def add_title_slide(prs, headline, subhead="Confidential Investor Presentation"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_brand_bar(slide)

    # Background accent block
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.2), Inches(13.33), Inches(1.3)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = THEME_SECONDARY
    bg.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.0), Inches(2.5))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = headline
    run.font.size = Pt(54)
    run.font.bold = True
    run.font.color.rgb = THEME_TEXT

    sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(4.0), Inches(11.0), Inches(1.0))
    stf = sub_box.text_frame
    stf.clear()
    sp = stf.paragraphs[0]
    sp.alignment = PP_ALIGN.LEFT
    srun = sp.add_run()
    srun.text = subhead
    srun.font.size = Pt(20)
    srun.font.color.rgb = THEME_MUTED

    # Logo placeholder
    logo = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), Inches(0.5), Inches(2.3), Inches(0.9))
    logo.fill.solid()
    logo.fill.fore_color.rgb = THEME_ACCENT
    logo.line.color.rgb = THEME_ACCENT
    ltf = logo.text_frame
    ltf.clear()
    lp = ltf.paragraphs[0]
    lp.alignment = PP_ALIGN.CENTER
    lrun = lp.add_run()
    lrun.text = "Your Logo"
    lrun.font.size = Pt(16)
    lrun.font.bold = True
    lrun.font.color.rgb = RGBColor(255, 255, 255)


def add_title_and_bullets(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    add_brand_bar(slide, title_text=title)
    slide.shapes.title.text = title
    set_text_format(slide.shapes.title.text_frame, font_size=34, bold=True)

    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()
    for idx, bullet in enumerate(bullets):
        p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
        p.text = bullet
        p.level = 0
        p.font.size = Pt(22)
        p.font.color.rgb = THEME_TEXT


def add_two_column(prs, title, left_points, right_points):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_brand_bar(slide, title_text=title)

    # Title text (visual)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.9))
    tf = title_box.text_frame
    tf.text = title
    set_text_format(tf, font_size=34, bold=True)

    # Left column
    left_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(5.9), Inches(4.5))
    ltf = left_box.text_frame
    ltf.clear()
    for i, item in enumerate(left_points):
        p = ltf.add_paragraph() if i > 0 else ltf.paragraphs[0]
        p.text = item
        p.level = 0
        p.font.size = Pt(22)
        p.font.color.rgb = THEME_TEXT

    # Right column
    right_box = slide.shapes.add_textbox(Inches(6.6), Inches(1.6), Inches(5.9), Inches(4.5))
    rtf = right_box.text_frame
    rtf.clear()
    for i, item in enumerate(right_points):
        p = rtf.add_paragraph() if i > 0 else rtf.paragraphs[0]
        p.text = item
        p.level = 0
        p.font.size = Pt(22)
        p.font.color.rgb = THEME_TEXT


def add_chart_slide(prs, title, categories, series_data, chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    add_brand_bar(slide, title_text=title)
    slide.shapes.title.text = title
    set_text_format(slide.shapes.title.text_frame, font_size=34, bold=True)

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for series_name, values in series_data:
        chart_data.add_series(series_name, values)

    x, y, cx, cy = Inches(1), Inches(1.7), Inches(11.3), Inches(4.8)
    chart_shape = slide.shapes.add_chart(chart_type, x, y, cx, cy, chart_data).chart
    chart_shape.has_legend = True
    if chart_shape.legend is not None:
        chart_shape.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart_shape.legend.include_in_layout = False
    chart_shape.has_title = False


def add_table_slide(prs, title, col_headers, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    add_brand_bar(slide, title_text=title)
    slide.shapes.title.text = title
    set_text_format(slide.shapes.title.text_frame, font_size=34, bold=True)

    rows_count = len(rows) + 1
    cols_count = len(col_headers)
    x, y, cx, cy = Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.0)
    table = slide.shapes.add_table(rows_count, cols_count, x, y, cx, cy).table

    for i, head in enumerate(col_headers):
        cell = table.cell(0, i)
        cell.text = head
        set_text_format(cell.text_frame, font_size=18, bold=True, color=THEME_PRIMARY)

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(val)
            set_text_format(cell.text_frame, font_size=16, bold=False)


def add_timeline(prs, title, milestones):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    add_brand_bar(slide, title_text=title)
    slide.shapes.title.text = title
    set_text_format(slide.shapes.title.text_frame, font_size=34, bold=True)

    start_x = 0.9
    y = 3.7
    spacing = (12.0 - start_x) / max(1, len(milestones) - 1)

    # Line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(y), Inches(11.7), Inches(0.1))
    line.fill.solid()
    line.fill.fore_color.rgb = THEME_MUTED
    line.line.fill.background()

    for i, (label, date) in enumerate(milestones):
        cx = start_x + i * spacing
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(y - 0.15), Inches(0.3), Inches(0.3))
        dot.fill.solid()
        dot.fill.fore_color.rgb = THEME_ACCENT if i == 0 else THEME_PRIMARY
        dot.line.fill.background()

        tbox = slide.shapes.add_textbox(Inches(cx - 0.7), Inches(y - 1.2), Inches(1.6), Inches(1.0))
        tf = tbox.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        run.font.size = Pt(16)
        run.font.color.rgb = THEME_TEXT

        dbox = slide.shapes.add_textbox(Inches(cx - 0.7), Inches(y + 0.2), Inches(1.6), Inches(0.6))
        dtf = dbox.text_frame
        dtf.clear()
        dp = dtf.paragraphs[0]
        dp.alignment = PP_ALIGN.CENTER
        drun = dp.add_run()
        drun.text = date
        drun.font.size = Pt(12)
        drun.font.color.rgb = THEME_MUTED


def build_pitch_deck(output_path: str, company_name: str = "Your Company", tagline: str = "One-line value proposition"):
    prs = Presentation()

    # 1 Title
    add_title_slide(prs, f"{company_name}", tagline)

    # 2 Problem
    add_title_and_bullets(prs, "Problem", [
        "[Who] experiences [pain] due to [root cause]",
        "The status quo is costly, slow, or error-prone",
        "Market shifts make the problem urgent now",
    ])

    # 3 Solution
    add_title_and_bullets(prs, "Solution", [
        "We provide [product] that delivers [key outcome]",
        "10x improvement in [metric] via [secret sauce]",
        "Simple, secure, and scalable by design",
    ])

    # 4 Market
    add_chart_slide(
        prs,
        "Market Size",
        categories=["Niche", "Beachhead", "Adjacent", "Total"],
        series_data=[("TAM/SAM/SOM (illustrative)", [50, 200, 600, 1200])],
        chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED,
    )

    # 5 Product
    add_two_column(
        prs,
        "Product",
        ["Core capability A", "Core capability B", "Core capability C"],
        ["Outcome 1", "Outcome 2", "Outcome 3"],
    )

    # 6 Business Model
    add_title_and_bullets(prs, "Business Model", [
        "Pricing: [subscription/usage/hybrid]",
        "Average contract value: $[x]",
        "Gross margin: [y]%",
    ])

    # 7 Traction
    add_chart_slide(
        prs,
        "Traction",
        categories=["M1", "M2", "M3", "M4", "M5", "M6"],
        series_data=[("Revenue ($k)", [2, 4, 6, 9, 15, 24]), ("Customers", [1, 3, 5, 8, 12, 18])],
        chart_type=XL_CHART_TYPE.LINE_MARKERS,
    )

    # 8 Competition
    add_table_slide(
        prs,
        "Competitive Landscape",
        ["Competitor", "Segment", "Strength", "Weakness"],
        [
            ["Incumbent A", "Enterprise", "Brand, reach", "Slow innovation"],
            ["Startup B", "SMB", "Usability", "Feature depth"],
            ["You", "Focused", "Speed, UX, insight", "Early-stage"],
        ],
    )

    # 9 Moat
    add_title_and_bullets(prs, "Defensibility (Moat)", [
        "Proprietary data from [source]",
        "Network effects across [participants]",
        "Workflow and switching costs",
    ])

    # 10 Go-To-Market
    add_title_and_bullets(prs, "Go-To-Market", [
        "Beachhead: [ICP / vertical]",
        "Channels: [direct, partners, PLG]",
        "Conversion: [top motion + key KPIs]",
    ])

    # 11 Roadmap
    add_timeline(prs, "Roadmap", [
        ("MVP", "Q1"),
        ("GA", "Q2"),
        ("Scale", "Q3"),
        ("New Verticals", "Q4"),
        ("International", "Q1 next"),
    ])

    # 12 Unit Economics
    add_chart_slide(
        prs,
        "Unit Economics",
        categories=["CAC", "LTV", "Payback (mo)", "Gross Margin %"],
        series_data=[("Metrics (illustrative)", [300, 2100, 6, 75])],
        chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED,
    )

    # 13 Financials
    add_chart_slide(
        prs,
        "3-Year Projections",
        categories=["Y1", "Y2", "Y3"],
        series_data=[("Revenue ($k)", [250, 1200, 3500]), ("OpEx ($k)", [600, 1400, 2200])],
        chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED,
    )

    # 14 Team
    add_title_and_bullets(prs, "Team", [
        "Founder A — ex-[Company], [domain] expert",
        "Founder B — ex-[Company], [tech] lead",
        "Advisors: [names]"
    ])

    # 15 Ask
    add_title_and_bullets(prs, "The Ask", [
        "Raising $[X] seed to reach [milestone]",
        "Use of funds: [key buckets]",
        "Runway: [N] months",
    ])

    # 16 Use of Funds
    add_table_slide(
        prs,
        "Use of Funds (Illustrative)",
        ["Category", "%", "$"],
        [["Engineering", "45%", "$450k"], ["Go-To-Market", "35%", "$350k"], ["G&A", "20%", "$200k"]],
    )

    # 17 Milestones
    add_timeline(prs, "Milestones", [
        ("Close round", "M0"),
        ("Hire core team", "M1-2"),
        ("GA + 10 design partners", "M3-4"),
        ("$100k MRR", "M9"),
        ("Series A ready", "M12"),
    ])

    # 18 Closing
    add_title_and_bullets(prs, "Thank You", [
        "Contact: founder@yourcompany.com",
        "Data room available upon request",
        "Appendix follows (optional)",
    ])

    prs.save(output_path)


if __name__ == "__main__":
    build_pitch_deck("/workspace/Investor_Pitch_Deck.pptx")

